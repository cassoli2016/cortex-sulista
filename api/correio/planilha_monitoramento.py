# -*- coding: utf-8 -*-
"""A planilha anexa ao monitoramento de cliente — no formato da torre.

O cliente recebe há meses uma planilha com estas colunas, nesta ordem, com o
cabeçalho em bege, uma faixa salmão por fluxo (mercadoria → destino) e a
situação escrita embaixo de cada carga. Mudar o formato junto com a
automação obrigaria quem lê a reaprender onde procurar no mesmo dia em que
precisa confiar numa fonte nova. Então o formato é o de sempre; o que muda é
de onde vêm os números.

A coluna MOTORISTA traz só o primeiro nome, como a torre escreve (ver o
cabeçalho de `api/correio/monitoramento`).

UMA DIFERENÇA DE PROPÓSITO:

- **Valores, não fórmulas.** A planilha da torre calcula TOTAL = saída −
  chegada na própria célula, e com a saída em branco a fórmula devolve um
  número negativo que o Excel mostra como "#####". Aqui o total vem pronto, e
  só existe com as DUAS pontas registradas: o veículo que ainda está no
  cliente fica com a saída, o total e as horas paradas em branco, como na
  planilha da torre (ver `monitoramento.dados`, sobre o relógio correndo).

Datas são datas e durações são durações (`[h]:mm`), não texto: quem recebe
soma horas paradas do mês em cima desta planilha, e texto não soma.
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# As cores da planilha da torre, lidas do arquivo que o cliente recebe hoje.
BEGE = "FFDDD9C3"          # cabeçalho
SALMAO = "FFFCE4D6"        # faixa do fluxo
SITUACAO = "FFFAE2D5"      # coluna da situação
AZUL = "FFC1E4F5"          # horas paradas

COLUNAS = [("FROTA", 16), ("PLACAS TRAÇÃO/CARRETA", 20), ("MOTORISTA", 13),
           ("JANELA CARREGAMENTO", 18), ("JANELA ENTREGA", 18),
           ("DATA E HORA CHEGADA NO CLIENTE", 20), ("DATA E HORA SAÍDA DO CLIENTE", 20),
           ("TOTAL", 9), ("CARÊNCIA", 10), ("TOTAL DE HORAS PARADAS", 13),
           ("NÚMERO DO PEDIDO", 12), ("CVA", 13), ("STATUS DO VEÍCULO", 62)]
N = len(COLUNAS)
COL_SIT = N                  # a última
COL_PARADAS = next(i for i, (n, _) in enumerate(COLUNAS, start=1)
                   if n == "TOTAL DE HORAS PARADAS")
DATA = "dd/mm/yyyy hh:mm"
DURACAO = "[h]:mm"

_fino = Side(style="thin", color="FFBFBFBF")
_BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


def _dt(txt):
    try:
        return datetime.strptime(str(txt)[:16], "%Y-%m-%d %H:%M") if txt else None
    except ValueError:
        return None


def _dur(h):
    return timedelta(hours=float(h)) if h is not None else None


def _cel(ws, lin, col, valor, *, fmt=None, fill=None, bold=False):
    c = ws.cell(row=lin, column=col, value=valor)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if bold:
        c.font = Font(bold=True)
    c.border = _BORDA
    c.alignment = Alignment(vertical="center", wrap_text=(col == COL_SIT))
    return c


def _situacao(c: dict) -> str:
    from api.correio.monitoramento import hm
    txt = c["marco"] + (f" · {_dt(c['marco_em']):%d/%m %H:%M}" if _dt(c["marco_em"]) else "")
    txt += f" ({c['marco_fonte']}" + (f", {c['marco_onde']}" if c["marco_onde"] else "") + ")"
    if c["chegada"] and not c["saida"]:
        txt += " · fim da descarga não registrado"
    elif c["permanencia_nd"]:
        txt += " · permanência acima de 24 h: n/d"
    elif c["paradas_h"]:
        txt += f" · {hm(c['paradas_h'])} além da carência"
    if c["eta"] and _dt(c["eta"]):
        txt += f" · previsão de chegada {_dt(c['eta']):%d/%m %H:%M} (histórico da rota)"
    return txt


def gerar(d: dict) -> bytes:
    """O .xlsx de uma rodada, a partir do mesmo `dados()` que monta o e-mail
    — o anexo e o corpo da mensagem não podem discordar."""
    agora = _dt(d["agora"]) or datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = f"{agora:%d.%m}"          # como as abas da torre: "11.09"

    for i, (rot, larg) in enumerate(COLUNAS, start=1):
        _cel(ws, 1, i, rot, fill=BEGE, bold=True)
        ws.cell(row=1, column=i).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    lin = 2
    for g in d["grupos"]:
        ws.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=N)
        _cel(ws, lin, 1, f"{g['mercadoria']} → {g['destinatario']}", fill=SALMAO, bold=True)
        lin += 1
        for c in g["cargas"]:
            vals = [c["frota"] or None, c["placas"] or None, c["motorista"] or None,
                    _dt(c["janela_carga"]), _dt(c["janela_entrega"]),
                    _dt(c["chegada"]), _dt(c["saida"]),
                    ("n/d" if c["permanencia_nd"] else _dur(c["permanencia_h"])),
                    _dur(c["carencia_h"]), _dur(c["paradas_h"]),
                    c["coleta"], c["cva"] or None, _situacao(c)]
            fmts = [None, None, None, DATA, DATA, DATA, DATA, DURACAO, DURACAO,
                    DURACAO, None, None, None]
            assert len(vals) == len(fmts) == N
            for i, (v, f) in enumerate(zip(vals, fmts), start=1):
                _cel(ws, lin, i, v, fmt=f,
                     fill=(SALMAO if i <= 2 else AZUL if i == COL_PARADAS
                           else SITUACAO if i == COL_SIT else None))
            lin += 1
            # O HISTÓRICO embaixo, uma linha por registro — é o que a torre
            # escreve à mão na coluna de situação ("Pátio carregado!!", "Em
            # viagem…", "Aguardando descarga…").
            for h in c["historico"]:
                q = _dt(h["quando"])
                _cel(ws, lin, COL_SIT,
                     f"{q:%d/%m %H:%M} — {h['rotulo']} ({h['fonte']})" if q
                     else f"{h['rotulo']} ({h['fonte']})", fill=SITUACAO)
                lin += 1
        lin += 1

    if not d["grupos"]:
        _cel(ws, lin, 1, "Nenhuma carga no dia.")
        lin += 2
    ws.cell(row=lin, column=1,
            value=f"Posição das {agora:%H:%M} de {agora:%d/%m/%Y} · fonte: {d['fonte']}"
            ).font = Font(italic=True, color="FF6B7580")
    ws.cell(row=lin + 1, column=1,
            value="Horas paradas = permanência no cliente além da carência do contrato "
                  "para a mercadoria, contada quando a chegada e o fim da descarga "
                  "estão registrados. Permanência acima de 24 h aparece como n/d."
            ).font = Font(italic=True, color="FF6B7580")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
