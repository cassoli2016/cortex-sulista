"""A planilha que o cliente recebe — no layout DELE.

"Cada cliente recebe de um jeito" é, na prática, três coisas: quais colunas,
em que ordem e com que nome. As três moram no perfil (`colunas`); aqui mora só
o CATÁLOGO do que existe para escolher e a escrita do arquivo.

O catálogo é a fronteira: um campo que não está nele não pode ir para a
planilha, por mais que o perfil peça. Isso impede que alguém "acrescente uma
coluna" que puxe dado que a tela não mostrou e ninguém conferiu.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

# campo -> (título padrão, tipo). O TIPO decide o formato da célula:
#   texto · numero · data_hora · duracao (segundos) · horas (freetime) · dinheiro
CATALOGO: dict[str, tuple[str, str]] = {
    "frota_cavalo":       ("Frota Cavalo", "texto"),
    "placa_cavalo":       ("Placa Cavalo", "texto"),
    "frota_carreta":      ("Frota Carreta", "texto"),
    "placa_carreta":      ("Placa Carreta", "texto"),
    "filial":             ("Filial", "numero"),
    "coleta":             ("Coleta", "numero"),
    "pedido":             ("Pedido (inteiro)", "texto"),
    "pedido_num":         ("Pedido", "texto"),
    "referencia":         ("Referência do cliente", "texto"),
    "mercadoria":         ("Mercadoria", "texto"),
    "origem":             ("Origem", "texto"),
    "destino":            ("Destino", "texto"),
    "cidade_origem":      ("Cidade de origem", "texto"),
    "cidade_destino":     ("Cidade de destino", "texto"),
    "ctes":               ("CT-e", "texto"),
    "carga_janela":       ("Janela Carregamento", "data_hora"),
    "carga_chegada":      ("Chegada Carregamento", "data_hora"),
    "carga_saida":        ("Saída Carregamento", "data_hora"),
    "carga_tempo":        ("Tempo Carregamento", "duracao"),
    "carga_freetime":     ("Freetime Carregamento", "horas"),
    "carga_valor_hora":   ("Valor da Hora", "dinheiro"),
    "carga_excedente":    ("Tempo Excedido Carregamento", "duracao"),
    "carga_valor":        ("Total Carregamento", "dinheiro"),
    "descarga_janela":    ("Janela Descarga", "data_hora"),
    "descarga_chegada":   ("Chegada Descarga", "data_hora"),
    "descarga_saida":     ("Saída Descarga", "data_hora"),
    "descarga_tempo":     ("Tempo Descarga", "duracao"),
    "descarga_freetime":  ("Freetime Descarga", "horas"),
    "descarga_valor_hora": ("Valor da Hora", "dinheiro"),
    "descarga_excedente": ("Tempo Excedido Descarga", "duracao"),
    "descarga_valor":     ("Total Descarga", "dinheiro"),
    "valor_total":        ("Total", "dinheiro"),
}

#: O layout de quem ainda não escolheu o seu.
COLUNAS_PADRAO = [
    "frota_cavalo", "placa_cavalo", "placa_carreta", "coleta", "pedido_num",
    "mercadoria", "origem", "destino",
    "carga_janela", "carga_chegada", "carga_saida", "carga_tempo",
    "descarga_janela", "descarga_chegada", "descarga_saida", "descarga_tempo",
    "ctes", "carga_freetime", "carga_excedente", "carga_valor",
    "descarga_freetime", "descarga_excedente", "descarga_valor", "valor_total",
]

# O que vai na célula, a partir da linha que a TELA mostra — a mesma linha,
# nunca uma leitura paralela: a planilha e a tela não podem discordar.
_PERNA = {"janela": "janela", "chegada": "chegada", "saida": "saida",
          "tempo": "tempo_s", "freetime": "freetime_h", "valor_hora": "valor_h",
          "excedente": "cobrado_s", "valor": "valor"}


def valor_do_campo(linha: dict, campo: str):
    if campo == "valor_total":
        return linha.get("valor")
    for perna in ("carga", "descarga"):
        if campo.startswith(perna + "_"):
            resto = campo[len(perna) + 1:]
            if resto in _PERNA:
                return (linha.get(perna) or {}).get(_PERNA[resto])
    return linha.get(campo)


def _dt(v):
    if v is None or isinstance(v, datetime):
        return v
    return datetime.fromisoformat(str(v).replace(" ", "T"))


def _hhmm(segundos: float) -> str:
    m = int(round(abs(segundos) / 60))
    return "%s%02d:%02d" % ("-" if segundos < 0 else "", m // 60, m % 60)


def nome_do_arquivo(modelo: str | None, contexto: dict) -> str:
    """O nome do arquivo, pelo modelo do perfil com `{{marcadores}}`.

    Substituição por regex de `{{nome_simples}}`, NUNCA `str.format` sobre
    texto escrito por usuário (regra da casa: `format` alcança atributo de
    objeto). Marcador desconhecido fica como está — e aparece no nome, que é
    onde quem escreveu vai perceber.
    """
    modelo = (modelo or "").strip() or "Horas paradas - {{cliente}} - {{de}} a {{ate}}"

    def troca(m):
        return str(contexto.get(m.group(1), m.group(0)))
    nome = re.sub(r"\{\{\s*([a-z_]+)\s*\}\}", troca, modelo)
    nome = re.sub(r'[\\/:*?"<>|]+', "-", nome).strip(" .") or "horas-paradas"
    return nome[:120] + ".xlsx"


def gerar(linhas: list[dict], colunas: list[dict], aba: str | None = None) -> bytes:
    """O `.xlsx`. `colunas` = [{campo, titulo}], na ordem do cliente.

    Só entram as linhas INCLUÍDAS — a excluída à mão continua na tela, com o
    motivo, e não vai para o cliente. A última linha soma as colunas de
    dinheiro, como na planilha que ele recebia.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\\/*?:\[\]]", "-", aba or "") or "Horas paradas")[:31]
    cols = [c for c in colunas if c.get("campo") in CATALOGO]

    cab = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="942821")
    for j, c in enumerate(cols, start=1):
        cel = ws.cell(row=1, column=j, value=c.get("titulo") or CATALOGO[c["campo"]][0])
        cel.font, cel.fill = cab, fundo
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = 14 if CATALOGO[c["campo"]][1] != "texto" else 20

    incluidas = [ln for ln in linhas if ln.get("incluida", True)]
    for i, ln in enumerate(incluidas, start=2):
        for j, c in enumerate(cols, start=1):
            tipo = CATALOGO[c["campo"]][1]
            v = valor_do_campo(ln, c["campo"])
            cel = ws.cell(row=i, column=j)
            if v is None or v == "":
                continue
            if tipo == "data_hora":
                cel.value = _dt(v)
                cel.number_format = "dd/mm/yyyy hh:mm"
            elif tipo == "duracao":
                # Excel não desenha hora negativa ("####"); o tempo negativo é
                # FATO (saiu antes da janela) e vai como texto legível.
                if v < 0:
                    cel.value = _hhmm(v)
                    cel.alignment = Alignment(horizontal="right")
                else:
                    cel.value = v / 86400
                    cel.number_format = "[h]:mm"
            elif tipo == "horas":
                cel.value = v / 24
                cel.number_format = "[h]:mm"
            elif tipo == "dinheiro":
                cel.value = float(v)
                cel.number_format = '"R$" #,##0.00'
            elif tipo == "numero":
                cel.value = v
            else:
                cel.value = str(v)

    if incluidas:
        fim = len(incluidas) + 2
        for j, c in enumerate(cols, start=1):
            if CATALOGO[c["campo"]][1] == "dinheiro" and "valor_hora" not in c["campo"]:
                total = sum(float(valor_do_campo(ln, c["campo"]) or 0) for ln in incluidas)
                cel = ws.cell(row=fim, column=j, value=round(total, 2))
                cel.number_format = '"R$" #,##0.00'
                cel.font = Font(bold=True)

    ws.freeze_panes = "A2"
    if cols:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), max(1, len(incluidas) + 1))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
