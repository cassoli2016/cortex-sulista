"""O controle de um cliente num período: ERP (dublê) + ajustes + regra."""
from __future__ import annotations

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from api.horas_paradas import cadastro, fonte, servico


def dt(s):
    return datetime.fromisoformat(s)


def _col(numero, **kw):
    base = {"grupo": 1, "empresa": 1, "filial": 20, "unidade": 1,
            "diferenciadornumero": 0, "serie": 1, "numero": numero,
            "emissao": dt("2026-09-08 12:00"), "pedido": "6100000001",
            "mercadoria": "PECAS", "placa_cavalo": "AAA0A00", "frota_cavalo": "A1",
            "placa_carreta": "BBB0B00", "frota_carreta": None,
            "origem": "PLANTA ORIGEM", "destino": "PLANTA DESTINO",
            "destinatario_codigo": "111", "cidade_origem": "X/SP", "cidade_destino": "Y/RJ",
            "carga_janela": dt("2026-09-09 10:00"), "carga_chegada": dt("2026-09-09 10:00"),
            "carga_saida": dt("2026-09-09 11:00"),
            "descarga_janela": dt("2026-09-09 15:00"), "descarga_chegada": dt("2026-09-09 15:00"),
            "descarga_saida": dt("2026-09-09 16:00"),
            "paradas": 0, "repeticoes": 1, "ctes": "100001"}
    base.update(kw)
    return base


CONTRATO = [{"filial": 20, "mercadoria": "", "ft_carga_h": 3.0, "ft_descarga_h": 3.0,
             "valor_coleta": 100.0, "valor_entrega": 100.0,
             "dtinicio": dt("2024-08-01 00:00"), "dtfim": None, "ativoinativo": 1}]

CARGAS = [
    _col(1, descarga_saida=dt("2026-09-09 18:30")),                  # 30 min excedidos
    _col(2),                                                         # dentro do freetime
    _col(3, descarga_chegada=None, descarga_saida=None),             # em aberto
    _col(4, carga_janela=dt("2026-09-01 10:00"), carga_chegada=dt("2026-09-01 10:00"),
         carga_saida=dt("2026-09-01 11:00"), descarga_janela=dt("2026-09-01 15:00"),
         descarga_chegada=dt("2026-09-01 15:00"),
         descarga_saida=dt("2026-09-01 16:00")),                     # outra semana
]


@pytest.fixture
def perfil(esquema_pg, monkeypatch):
    monkeypatch.setattr(fonte, "cargas", lambda cli, de, ate: {
        "cargas": [dict(c) for c in CARGAS], "contrato": [dict(x) for x in CONTRATO],
        "lido_em": "2026-09-12T10:00:00"})
    p = cadastro.criar_perfil(8, "CLIENTE A", "a@x", esquema=esquema_pg)
    return p["id"], esquema_pg


def _montar(perfil, de="2026-09-07", ate="2026-09-13"):
    pid, esq = perfil
    return servico.montar(pid, de, ate, esquema=esq)


def test_o_periodo_se_decide_pelo_MARCO_do_perfil(perfil):
    d = _montar(perfil)
    assert [ln["coleta"] for ln in d["linhas"]] == [1, 2]
    assert [a["coleta"] for a in d["abertas"]] == [3]


def test_carga_em_aberto_NAO_soma(perfil):
    """Somar pela metade seria cobrar duas vezes quando ela fechar."""
    r = _montar(perfil)["resumo"]
    assert r["cargas"] == 2 and r["em_aberto"] == 1
    assert r["valor_total"] == 50.0 and r["com_excedente"] == 1


def test_ajuste_muda_a_conta_e_a_linha_MOSTRA_o_lado_do_ERP(perfil):
    pid, esq = perfil
    ch = fonte.chave(CARGAS[1])
    cadastro.ajustar(ch, "descarga_saida", "2026-09-09 19:00", "2026-09-09T16:00",
                     "fim de descarga apontado atrasado", "a@x", esquema=esq)
    d = _montar(perfil)
    ln = next(x for x in d["linhas"] if x["coleta"] == 2)
    assert ln["descarga"]["saida"] == "2026-09-09 19:00"
    assert ln["descarga"]["erp"] == {"saida": "2026-09-09 16:00"}
    assert ln["valor"] == 100.0 and ln["valor_erp"] == 0.0
    # O EFEITO DO AJUSTE APARECE NO TOPO, em reais
    assert d["resumo"]["ajustadas"] == 1 and d["resumo"]["efeito_ajustes"] == 100.0


def test_ajuste_que_FECHA_a_carga_em_aberto_a_traz_para_o_periodo(perfil):
    pid, esq = perfil
    ch = fonte.chave(CARGAS[2])
    cadastro.ajustar(ch, "descarga_chegada", "2026-09-09 15:00", None, "portaria", "a@x", esquema=esq)
    cadastro.ajustar(ch, "descarga_saida", "2026-09-09 16:30", None, "portaria", "a@x", esquema=esq)
    d = _montar(perfil)
    assert 3 in [ln["coleta"] for ln in d["linhas"]] and not d["abertas"]


def test_carga_EXCLUIDA_fica_na_tela_mas_nao_soma_nem_vai_na_planilha(perfil):
    pid, esq = perfil
    cadastro.ajustar(fonte.chave(CARGAS[0]), "incluir", "nao", None,
                     "cliente cobre esta estadia em contrato próprio", "a@x", esquema=esq)
    d = _montar(perfil)
    ln = next(x for x in d["linhas"] if x["coleta"] == 1)
    assert ln["incluida"] is False
    assert d["resumo"]["valor_total"] == 0.0 and d["resumo"]["excluidas"] == 1
    _, xb = servico.exportar(pid, "2026-09-07", "2026-09-13", esquema=esq)
    ws = load_workbook(io.BytesIO(xb)).active
    coletas = [ws.cell(r, c).value for r in range(2, ws.max_row + 1)
               for c in range(1, ws.max_column + 1) if ws.cell(1, c).value == "Coleta"]
    assert 1 not in coletas and 2 in coletas


def test_recorte_pela_JANELA_DE_CARGA_muda_quem_entra(perfil):
    pid, esq = perfil
    cadastro.salvar_config(pid, {"recorte": "janela_carga"}, "a@x", esquema=esq)
    d = _montar(perfil)
    assert [ln["coleta"] for ln in d["linhas"]] == [1, 2, 3]


def test_clausula_vale_NA_DATA_DA_CARGA_e_nao_hoje(perfil, monkeypatch):
    """Cobrar a semana passada com a cláusula que entrou ontem seria cobrar
    retroativo."""
    novo = dict(CONTRATO[0], ft_descarga_h=1.0, dtinicio=dt("2026-09-10 00:00"))
    velho = dict(CONTRATO[0], dtfim=dt("2026-09-09 23:59"))
    monkeypatch.setattr(fonte, "cargas", lambda cli, de, ate: {
        "cargas": [dict(CARGAS[0])], "contrato": [novo, velho], "lido_em": "x"})
    ln = _montar(perfil)["linhas"][0]
    assert ln["descarga"]["freetime_h"] == 3.0


def test_periodo_invertido_e_longo_demais_sao_recusados(perfil):
    with pytest.raises(cadastro.Recusa):
        _montar(perfil, "2026-09-13", "2026-09-07")
    with pytest.raises(cadastro.Recusa):
        _montar(perfil, "2026-01-01", "2026-09-07")


def test_perfil_inexistente(esquema_pg):
    with pytest.raises(servico.PerfilNaoExiste):
        servico.montar(999, "2026-09-07", "2026-09-13", esquema=esquema_pg)


@pytest.mark.parametrize("pedido, esperado", [
    ("6100000123ABC0000456", ("6100000123", "ABC0000456")),
    ("6100000124 OPERACAO", ("6100000124", "")),
    ("6100000125", ("6100000125", "")),
    ("SEM ORDEM DE FRETE", ("SEM ORDEM DE FRETE", "")),
    ("", ("", "")),
])
def test_pedido_e_o_complemento_colado_nele(pedido, esperado):
    assert servico.pedido_partes(pedido) == esperado


def test_linha_avisa_quando_a_coleta_tem_mais_de_um_apontamento(perfil, monkeypatch):
    monkeypatch.setattr(fonte, "cargas", lambda cli, de, ate: {
        "cargas": [dict(CARGAS[0], repeticoes=2, paradas=3)], "contrato": CONTRATO,
        "lido_em": "x"})
    avisos = " ".join(_montar(perfil)["linhas"][0]["avisos"])
    assert "mais de um apontamento" in avisos and "3 paradas" in avisos
