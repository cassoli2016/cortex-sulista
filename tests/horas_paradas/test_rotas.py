"""As rotas de Horas Paradas com sessão de verdade.

O que só se prova aqui: recusa é 4xx (o Cloudflare troca o corpo dos 5xx),
sem sessão é 401, a escrita entra no `audit_log`, e a planilha baixada sai
das MESMAS linhas que a rota da tela devolve.
"""
from __future__ import annotations

import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import auth
from api.horas_paradas import cadastro, fonte

SENHA = "senha-de-teste-123"


def dt(s):
    return datetime.fromisoformat(s)


CARGA = {"grupo": 1, "empresa": 1, "filial": 20, "unidade": 1,
         "diferenciadornumero": 0, "serie": 1, "numero": 7,
         "emissao": dt("2026-09-08 12:00"), "pedido": "6100000007CIF0000007",
         "mercadoria": "PECAS", "placa_cavalo": "AAA0A00", "frota_cavalo": "A1",
         "placa_carreta": "", "frota_carreta": None, "origem": "O", "destino": "D",
         "destinatario_codigo": "111", "cidade_origem": "", "cidade_destino": "",
         "carga_janela": dt("2026-09-09 10:00"), "carga_chegada": dt("2026-09-09 10:00"),
         "carga_saida": dt("2026-09-09 11:00"),
         "descarga_janela": dt("2026-09-09 15:00"), "descarga_chegada": dt("2026-09-09 15:00"),
         "descarga_saida": dt("2026-09-09 19:00"),
         "paradas": 0, "repeticoes": 1, "ctes": "1"}
CONTRATO = [{"filial": 20, "mercadoria": "", "ft_carga_h": 3.0, "ft_descarga_h": 3.0,
             "valor_coleta": 100.0, "valor_entrega": 100.0,
             "dtinicio": None, "dtfim": None, "ativoinativo": 1}]


@pytest.fixture
def cliente(esquema_pg, monkeypatch):
    from api.main import app
    monkeypatch.setattr(auth, "ESQUEMA", esquema_pg)
    monkeypatch.setattr(cadastro, "ESQUEMA", esquema_pg)
    monkeypatch.setattr(fonte, "cargas", lambda cli, de, ate: {
        "cargas": [dict(CARGA)], "contrato": [dict(x) for x in CONTRATO],
        "lido_em": "2026-09-12T10:00:00"})
    auth.init_db()
    with auth._conn() as c:
        perfil = c.execute(
            "SELECT id FROM perfis WHERE admin=1 ORDER BY id LIMIT 1").fetchone()
        c.execute(
            """INSERT INTO usuarios(nome, email, senha_hash, perfil_id, ativo,
                                    deve_trocar_senha, criado_em)
               VALUES('Chefe','chefe@sulista.com.br',%s,%s,1,0,%s)""",
            (auth._ph.hash(SENHA), perfil["id"], auth._agora()))
    cli = TestClient(app)
    r = cli.post("/api/auth/login", json={"email": "chefe@sulista.com.br", "senha": SENHA})
    assert r.status_code == 200, r.text
    return cli


@pytest.mark.parametrize("caminho", ["/api/operacao/horas-paradas?perfil=1",
                                     "/api/operacao/horas-paradas/perfis"])
def test_sem_sessao_e_401(caminho):
    from api.main import app
    assert TestClient(app).get(caminho).status_code == 401


def test_rota_mapeada_para_a_tela_hp():
    """Fora de `ROTA_TELAS` o middleware (fail-closed) daria 403 a quem tem a
    tela — e o defeito só apareceria para um usuário não-admin."""
    alvo = [t for p, t in auth.ROTA_TELAS if "/api/operacao/horas-paradas/ajuste".startswith(p)]
    assert alvo and "hp" in alvo[0]
    assert "hp" in auth.TELAS


def test_ciclo_perfil_regra_ajuste_planilha(cliente):
    r = cliente.post("/api/operacao/horas-paradas/perfis",
                     json={"cliente_codigo": 8, "cliente_nome": "CLIENTE A"})
    assert r.status_code == 200, r.text
    pid = r.json()["perfil"]["id"]

    r = cliente.post("/api/operacao/horas-paradas/perfis/salvar",
                     json={"id": pid, "config": {"arredondamento_min": 60}})
    assert r.status_code == 200, r.text

    d = cliente.get("/api/operacao/horas-paradas",
                    params={"perfil": pid, "de": "2026-09-07", "ate": "2026-09-13"}).json()
    assert d["resumo"]["valor_total"] == 100.0          # 1h de descarga excedida
    ln = d["linhas"][0]
    assert ln["pedido_num"] == "6100000007" and ln["referencia"] == "CIF0000007"

    r = cliente.post("/api/operacao/horas-paradas/ajuste", json={
        "chave": ln["chave"], "campo": "descarga_saida", "valor": "2026-09-09 20:10",
        "valor_erp": "2026-09-09 19:00", "motivo": "fim de descarga conferido na portaria"})
    assert r.status_code == 200, r.text
    d = cliente.get("/api/operacao/horas-paradas",
                    params={"perfil": pid, "de": "2026-09-07", "ate": "2026-09-13"}).json()
    assert d["resumo"]["valor_total"] == 300.0          # 2h10 → 3h (arredonda p/ cima)
    assert d["resumo"]["efeito_ajustes"] == 200.0

    x = cliente.get("/api/operacao/horas-paradas/planilha",
                    params={"perfil": pid, "de": "2026-09-07", "ate": "2026-09-13"})
    assert x.status_code == 200
    assert "spreadsheetml" in x.headers["content-type"]
    assert "attachment" in x.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(x.content)).active
    ultima = [c.value for c in ws[ws.max_row]]
    assert 300.0 in ultima, "o total da planilha tem de ser o da tela: %s" % ultima

    with auth._conn() as c:
        acoes = [r["acao"] for r in c.execute(
            "SELECT acao FROM audit_log WHERE acao LIKE 'hp_%%' ORDER BY id").fetchall()]
    assert acoes == ["hp_perfil_criar", "hp_perfil_salvar", "hp_ajuste"]


def test_recusas_sao_4xx_legiveis(cliente):
    pid = cliente.post("/api/operacao/horas-paradas/perfis",
                       json={"cliente_codigo": 8, "cliente_nome": "A"}).json()["perfil"]["id"]
    r = cliente.post("/api/operacao/horas-paradas/perfis",
                     json={"cliente_codigo": 8, "cliente_nome": "A"})
    assert r.status_code == 409 and "já tem perfil" in r.json()["mensagem"]
    r = cliente.post("/api/operacao/horas-paradas/perfis/salvar",
                     json={"id": pid, "config": {"inicio_carga": "amanha"}})
    assert r.status_code == 409 and "Início do relógio" in r.json()["mensagem"]
    r = cliente.post("/api/operacao/horas-paradas/ajuste", json={
        "chave": "1|1|20|1|0|1|7", "campo": "carga_saida", "valor": "2026-09-09 12:00",
        "motivo": "x"})
    assert r.status_code == 409 and "motivo" in r.json()["mensagem"]
    r = cliente.get("/api/operacao/horas-paradas", params={"perfil": 999})
    assert r.status_code == 404
    r = cliente.get("/api/operacao/horas-paradas",
                    params={"perfil": pid, "de": "ontem", "ate": "hoje"})
    assert r.status_code == 422
