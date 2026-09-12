"""A planilha que vai para o cliente: colunas DELE, na ordem DELE."""
from __future__ import annotations

import io
from datetime import timedelta

from openpyxl import load_workbook

from api.horas_paradas import planilha


def _linha(coleta, valor_c, valor_d, incluida=True, tempo=3600.0):
    return {
        "coleta": coleta, "placa_cavalo": "AAA0A00", "mercadoria": "PECAS",
        "incluida": incluida, "valor": valor_c + valor_d,
        "carga": {"janela": "2026-09-09 10:00", "chegada": "2026-09-09 09:00",
                  "saida": "2026-09-09 11:00", "tempo_s": tempo, "freetime_h": 3.0,
                  "valor_h": 100.0, "cobrado_s": 0.0, "valor": valor_c},
        "descarga": {"janela": None, "chegada": None, "saida": None, "tempo_s": None,
                     "freetime_h": 6.5, "valor_h": 100.0, "cobrado_s": 5400.0,
                     "valor": valor_d},
    }


def _ler(xb):
    ws = load_workbook(io.BytesIO(xb)).active
    return ws, [[c.value for c in r] for r in ws.iter_rows()]


COLS = [{"campo": "coleta", "titulo": "Coleta"},
        {"campo": "placa_cavalo", "titulo": "Placa"},
        {"campo": "carga_tempo", "titulo": "Tempo Carregamento"},
        {"campo": "descarga_freetime", "titulo": "Freetime Descarga"},
        {"campo": "carga_valor_hora", "titulo": "Valor da Hora"},
        {"campo": "carga_valor", "titulo": "Total Carregamento"},
        {"campo": "descarga_valor", "titulo": "Total Descarga"}]


def test_cabecalho_e_a_ordem_sao_os_do_perfil():
    _, rows = _ler(planilha.gerar([_linha(1, 0, 0)], COLS))
    assert rows[0] == [c["titulo"] for c in COLS]


def test_linha_EXCLUIDA_nao_vai_para_o_cliente():
    ws, rows = _ler(planilha.gerar([_linha(1, 10, 0), _linha(2, 20, 0, incluida=False)], COLS))
    assert [r[0] for r in rows[1:-1]] == [1]


def test_total_soma_so_DINHEIRO_e_nao_o_valor_da_hora():
    """O cliente recebia os totais embaixo das colunas de total; somar o valor
    da hora de cada linha daria um número sem sentido."""
    _, rows = _ler(planilha.gerar([_linha(1, 34.01, 0), _linha(2, 0, 150.0)], COLS))
    total = rows[-1]
    assert total[5] == 34.01 and total[6] == 150.0
    assert total[4] is None


def test_duracao_e_freetime_saem_como_TEMPO_do_excel():
    """Célula de TEMPO, não texto: o cliente soma a coluna. (O openpyxl lê de
    volta a célula `[h]:mm` como `timedelta` — é a prova de que é tempo.)"""
    ws, rows = _ler(planilha.gerar([_linha(1, 0, 0)], COLS))
    assert rows[1][2] == timedelta(hours=1) and ws.cell(2, 3).number_format == "[h]:mm"
    assert rows[1][3] == timedelta(hours=6, minutes=30)


def test_tempo_NEGATIVO_vai_como_texto_legivel():
    """O Excel desenha "####" para hora negativa; o fato (saiu antes da janela)
    tem de chegar lido."""
    _, rows = _ler(planilha.gerar([_linha(1, 0, 0, tempo=-300.0)], COLS))
    assert rows[1][2] == "-00:05"


def test_campo_fora_do_catalogo_NAO_entra():
    _, rows = _ler(planilha.gerar([_linha(1, 0, 0)], COLS + [{"campo": "senha", "titulo": "x"}]))
    assert "x" not in rows[0]


def test_nome_do_arquivo_troca_marcadores_por_regex_e_nao_por_format():
    ctx = {"cliente": "CLIENTE A", "semana": 37, "de": "07-09-2026", "ate": "13-09-2026"}
    assert (planilha.nome_do_arquivo("HP semana - Semana {{semana}} - {{de}}", ctx)
            == "HP semana - Semana 37 - 07-09-2026.xlsx")
    # `str.format` alcançaria atributo de objeto; aqui o texto fica como está
    assert "{0.__class__}" in planilha.nome_do_arquivo("x {0.__class__}", ctx)
    assert planilha.nome_do_arquivo("a/b:c", ctx) == "a-b-c.xlsx"
    assert planilha.nome_do_arquivo("", ctx).startswith("Horas paradas - CLIENTE A")
