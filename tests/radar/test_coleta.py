# -*- coding: utf-8 -*-
"""A coleta do Radar: idempotente, com cadência, e falha que não apaga."""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone

import openpyxl

from api import pglocal
from api.radar import coleta, fontes
from tests.radar.conftest import TomTom, ler

TODAS = {"anp", "brent", "dolar", "ptax", "rodovias", "frota",
         *[f"noticias_{t}" for t in fontes.TEMAS]}


def _n(esq, tabela, onde="", params=()) -> int:
    return pglocal.um(f"SELECT count(*) AS n FROM {tabela} {onde}", params,
                      esquema=esq)["n"]


def _coletar(esq, rede, tomtom, **kw):
    return coleta.coletar(esquema=esq, baixar=rede, consultar_tomtom=tomtom, **kw)


def test_a_primeira_passada_grava_todas_as_fontes(esquema_pg, rede, tomtom, relogio):
    r = _coletar(esquema_pg, rede, tomtom)
    assert set(r) == TODAS
    assert set(r.values()) == {"ok"}, r
    assert _n(esquema_pg, "rad_combustivel") == 12
    assert _n(esquema_pg, "rad_serie", "WHERE serie = 'brent'") == 45
    assert _n(esquema_pg, "rad_serie", "WHERE serie = 'dolar'") == 44
    assert _n(esquema_pg, "rad_serie", "WHERE serie = 'ptax'") == 51
    cot = pglocal.um("SELECT * FROM rad_cotacao WHERE serie = 'brent'", esquema=esquema_pg)
    assert float(cot["valor"]) == 103.68 and cot["fuso"] == "America/New_York"
    assert _n(esquema_pg, "rad_rodovia") == 8, "dois relevantes em cada um dos 4 corredores"
    assert _n(esquema_pg, "rad_noticia") == 3 * len(fontes.TEMAS)


def test_a_primeira_carga_pede_dois_anos_e_as_seguintes_um_mes(esquema_pg, rede, tomtom,
                                                               relogio, monkeypatch):
    _coletar(esquema_pg, rede, tomtom)
    assert any("BZ%3DF" in u and "range=2y" in u for u in rede.pedidos)
    monkeypatch.setattr(coleta, "PONTOS_CARGA_COMPLETA", 10)
    rede.pedidos.clear()
    _coletar(esquema_pg, rede, tomtom, forcar=True, so={"brent"})
    assert [u for u in rede.pedidos if "BZ%3DF" in u][0].count("range=1mo") == 1


def test_a_segunda_passada_respeita_a_cadencia_e_nao_baixa_nada(esquema_pg, rede,
                                                                 tomtom, relogio):
    _coletar(esquema_pg, rede, tomtom)
    pedidos, chamadas = len(rede.pedidos), tomtom.chamadas
    relogio.andar(minutes=5)
    r = _coletar(esquema_pg, rede, tomtom)
    assert set(r.values()) == {"no_prazo"}
    assert len(rede.pedidos) == pedidos and tomtom.chamadas == chamadas


def test_o_preco_de_agora_volta_a_cada_ciclo_de_dez_minutos(esquema_pg, rede, tomtom,
                                                            relogio):
    """A cadência do mercado é um pouco menor que o ciclo da thread: com ela
    igual ao ciclo, um atraso de milissegundos pularia um ciclo inteiro."""
    _coletar(esquema_pg, rede, tomtom)
    relogio.andar(seconds=coleta.CADENCIA_S["brent"] + 1)
    r = _coletar(esquema_pg, rede, tomtom)
    assert r["brent"] == "ok" and r["dolar"] == "ok"
    assert r["anp"] == "no_prazo"
    assert coleta.CADENCIA_S["brent"] < 600


def test_forcar_regrava_sem_duplicar(esquema_pg, rede, tomtom, relogio):
    _coletar(esquema_pg, rede, tomtom)
    antes = {t: _n(esquema_pg, t) for t in
             ("rad_combustivel", "rad_serie", "rad_noticia", "rad_rodovia", "rad_cotacao")}
    _coletar(esquema_pg, rede, tomtom, forcar=True)
    depois = {t: _n(esquema_pg, t) for t in antes}
    assert antes == depois


def test_fonte_fora_do_ar_NAO_APAGA_o_que_ja_estava(esquema_pg, rede, tomtom, relogio):
    _coletar(esquema_pg, rede, tomtom)
    rede.falhar.add("gov.br/anp")
    r = _coletar(esquema_pg, rede, tomtom, forcar=True)
    assert r["anp"] == "erro"
    assert r["brent"] == "ok", "uma fonte fora não derruba as outras"
    assert _n(esquema_pg, "rad_combustivel") == 12
    e = coleta.estado(esquema_pg)["anp"]
    assert e["ok"] is False and e["erro"] == "HTTPError 503"
    assert e["sucesso_em"] is not None and e["itens"] == 12, \
        "a falha não pode apagar o registro do último sucesso"


def test_a_falha_volta_em_meia_hora_e_nao_no_proximo_ciclo(esquema_pg, rede, tomtom,
                                                           relogio):
    rede.falhar.add("gov.br/anp")
    _coletar(esquema_pg, rede, tomtom)
    relogio.andar(minutes=10)
    assert _coletar(esquema_pg, rede, tomtom)["anp"] == "no_prazo"
    relogio.andar(minutes=21)
    assert _coletar(esquema_pg, rede, tomtom)["anp"] == "erro", "tentou de novo"


def test_preco_fora_da_faixa_fisica_e_recusado_e_nao_gravado(esquema_pg, rede, tomtom,
                                                             relogio):
    corpo = json.loads(ler("yahoo_brent_trecho.json"))
    corpo["chart"]["result"][0]["meta"]["regularMarketPrice"] = 1036.8
    rede.trocar["BZ%3DF"] = json.dumps(corpo).encode()
    r = _coletar(esquema_pg, rede, tomtom)
    assert r["brent"] == "erro"
    assert _n(esquema_pg, "rad_cotacao", "WHERE serie = 'brent'") == 0
    assert "faixa" in coleta.estado(esquema_pg)["brent"]["erro"]


def test_planilha_sem_diesel_e_ERRO_e_nao_sucesso_vazio(esquema_pg, rede, tomtom, relogio):
    wb = openpyxl.load_workbook(io.BytesIO(ler("anp_semanal_trecho.xlsx")))
    ws = wb.active
    cab = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "DATA INICIAL")
    ws.delete_rows(cab + 1, ws.max_row)
    buf = io.BytesIO()
    wb.save(buf)
    rede.trocar["gov.br/anp"] = buf.getvalue()
    r = _coletar(esquema_pg, rede, tomtom)
    assert r["anp"] == "erro"
    assert "formato inesperado" in coleta.estado(esquema_pg)["anp"]["erro"]
    assert _n(esquema_pg, "rad_combustivel") == 0


# ---------------------------------------------------------------- rodovias

def test_rodovias_com_todos_os_corredores_fora_mantem_o_retrato_velho(esquema_pg, rede,
                                                                       tomtom, relogio):
    _coletar(esquema_pg, rede, tomtom)
    r = _coletar(esquema_pg, rede, TomTom(falhar=True), forcar=True, so={"rodovias"})
    assert r["rodovias"] == "erro"
    assert _n(esquema_pg, "rad_rodovia") == 8


def test_rodovias_com_zero_bruto_e_leitura_SUSPEITA(esquema_pg, rede, relogio):
    """A Grande SP nunca fica sem uma obra registrada: zero nos quatro
    corredores é a API respondendo mal, não a estrada livre."""
    r = _coletar(esquema_pg, rede, TomTom(corpo={"incidents": []}), so={"rodovias"})
    assert r["rodovias"] == "erro"


def test_rodovias_com_zero_DEPOIS_do_filtro_e_retrato_vazio_verdadeiro(esquema_pg, rede,
                                                                       tomtom, relogio):
    _coletar(esquema_pg, rede, tomtom)
    so_rua = {"incidents": [i for i in tomtom.corpo["incidents"]
                            if not i["properties"]["roadNumbers"]]}
    r = _coletar(esquema_pg, rede, TomTom(corpo=so_rua), forcar=True, so={"rodovias"})
    assert r["rodovias"] == "ok"
    assert _n(esquema_pg, "rad_rodovia") == 0


def test_o_consumo_da_TomTom_e_contado_NO_SCHEMA_DO_TESTE(esquema_pg, rede, tomtom, relogio):
    """Sem o esquema passado adiante, o contador cairia no schema de produção."""
    _coletar(esquema_pg, rede, tomtom, so={"rodovias"})
    l = pglocal.um("SELECT chamadas FROM tt_chamadas WHERE recurso = 'radar_incidentes'",
                   esquema=esquema_pg)
    assert l and l["chamadas"] == 4


def test_sem_TomTom_a_fonte_nem_entra_no_plano(esquema_pg, rede, relogio, monkeypatch):
    """Recurso não contratado não é falha: sem esta regra a Saúde ficaria
    vermelha a cada 20 minutos por um motivo que ninguém precisa consertar."""
    from api.radar import rodovias
    monkeypatch.setattr(rodovias, "ativo", lambda: False)
    r = coleta.coletar(esquema=esquema_pg, baixar=rede)
    assert "rodovias" not in r
    assert "rodovias" not in coleta.estado(esquema_pg)


# ---------------------------------------------------------------- notícias

def test_noticia_velha_nao_entra_e_a_poda_tira_a_de_noventa_dias(esquema_pg, rede, tomtom,
                                                                  relogio):
    pglocal.executar(
        "INSERT INTO rad_noticia (tema, guid, titulo, link, publicada_em) "
        "VALUES ('diesel', 'velha', 'Manchete de junho', 'https://x.test/a', %s)",
        (datetime(2026, 6, 1, tzinfo=timezone.utc),), esquema=esquema_pg)
    relogio.andar(days=40)
    r = _coletar(esquema_pg, rede, tomtom, so={"noticias_diesel"})
    assert r["noticias_diesel"] == "ok", "busca sem manchete nova é possível, não é erro"
    assert _n(esquema_pg, "rad_noticia") == 0
    assert coleta.estado(esquema_pg)["noticias_diesel"]["descartados"] == 3
