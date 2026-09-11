# -*- coding: utf-8 -*-
"""Os leitores das fontes públicas, contra os corpos REAIS de 11/09/2026."""
from __future__ import annotations

import io
import json
import urllib.error
from datetime import date, datetime, timezone

import openpyxl
import pytest

from api.radar import fontes
from tests.radar.conftest import ler


# ------------------------------------------------------------------- ANP

def _anp_editada(editar) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(ler("anp_semanal_trecho.xlsx")))
    ws = wb.active
    linha_cab = next(r for r in range(1, ws.max_row + 1)
                     if ws.cell(r, 1).value == "DATA INICIAL")
    editar(ws, linha_cab)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _s10_da_ultima_semana(linhas):
    return next(l for l in linhas
                if l["produto"] == "diesel_s10" and l["semana_fim"] == date(2026, 9, 5))


def test_anp_le_SO_as_duas_linhas_de_diesel():
    """A planilha tem sete combustíveis; a casa roda a diesel. E "ÓLEO DIESEL"
    sem sufixo é o S500 — a própria nota do arquivo diz."""
    linhas = fontes.ler_anp(ler("anp_semanal_trecho.xlsx"))
    assert {l["produto"] for l in linhas} == {"diesel_s10", "diesel_s500"}
    assert len(linhas) == 12, "seis semanas × dois produtos"
    s10 = _s10_da_ultima_semana(linhas)
    assert (s10["preco_revenda"], s10["preco_min"], s10["preco_max"]) == (6.88, 5.63, 9.25)
    assert s10["postos"] == 3135
    assert s10["semana_inicio"] == date(2026, 8, 30)


def test_o_traco_da_planilha_e_AUSENCIA_e_nao_zero():
    """O preço de distribuição vem '-' desde ago/2020. Zero seria um preço."""
    s10 = _s10_da_ultima_semana(fontes.ler_anp(ler("anp_semanal_trecho.xlsx")))
    assert s10["preco_distribuicao"] is None


def test_anp_le_pelo_CABECALHO_e_nao_pela_posicao():
    """Uma coluna nova antes do PRODUTO não pode fazer o desvio-padrão virar
    preço médio."""
    def coluna_nova(ws, cab):
        ws.insert_cols(3)
        ws.cell(cab, 3).value = "COLUNA QUE A ANP INVENTOU"
    s10 = _s10_da_ultima_semana(fontes.ler_anp(_anp_editada(coluna_nova)))
    assert s10["preco_revenda"] == 6.88


def test_anp_sem_a_coluna_de_preco_e_RECUSA_nomeando_a_coluna():
    def renomeia(ws, cab):
        for c in ws[cab]:
            if c.value and fontes._norm(c.value) == "PRECO MEDIO REVENDA":
                c.value = "OUTRA COISA"
    with pytest.raises(fontes.FormatoInesperado, match="PRECO MEDIO REVENDA"):
        fontes.ler_anp(_anp_editada(renomeia))


def test_o_que_nao_e_planilha_e_recusa_e_nao_lista_vazia():
    with pytest.raises(fontes.FormatoInesperado):
        fontes.ler_anp(b"<!DOCTYPE html><html>manutencao</html>")


# ----------------------------------------------------------------- Yahoo

def test_o_preco_de_agora_carrega_a_hora_do_ULTIMO_NEGOCIO():
    """Num sábado a hora é de sexta — é ela, e não a da coleta, que a tela
    mostra ao lado do preço."""
    y = fontes.ler_yahoo(ler("yahoo_brent_trecho.json"))
    assert y["agora"]["valor"] == 103.68
    assert y["agora"]["momento"] == datetime.fromtimestamp(1789122402, timezone.utc)
    assert y["agora"]["simbolo"] == "BZ=F"
    assert len(y["pontos"]) == 45
    assert y["pontos"][-1][0] == date(2026, 9, 11)


def test_fechamento_nulo_sai_e_nao_vira_zero():
    """O corpo real do dólar tem um pregão com fechamento `null`."""
    corpo = json.loads(ler("yahoo_dolar_trecho.json"))
    nulos = corpo["chart"]["result"][0]["indicators"]["quote"][0]["close"].count(None)
    assert nulos == 1, "a amostra real mudou — este teste depende do nulo dela"
    y = fontes.ler_yahoo(ler("yahoo_dolar_trecho.json"))
    assert len(y["pontos"]) == 44
    assert all(v > 1 for _, v in y["pontos"])


def test_o_dia_da_barra_sai_no_FUSO_DA_BOLSA():
    corpo = json.loads(ler("yahoo_dolar_trecho.json"))
    r = corpo["chart"]["result"][0]
    from zoneinfo import ZoneInfo
    esperado = datetime.fromtimestamp(r["timestamp"][-1],
                                      ZoneInfo(r["meta"]["exchangeTimezoneName"])).date()
    assert fontes.ler_yahoo(ler("yahoo_dolar_trecho.json"))["pontos"][-1][0] == esperado


def test_a_recusa_do_proprio_Yahoo_e_formato_inesperado():
    corpo = (b'{"chart":{"result":null,"error":{"code":"Not Found",'
             b'"description":"No data found, symbol may be delisted"}}}')
    with pytest.raises(fontes.FormatoInesperado, match="Not Found"):
        fontes.ler_yahoo(corpo)


# ------------------------------------------------------------------ SGS

def test_a_ptax_sai_do_sgs_com_data_brasileira():
    pontos = fontes.ler_sgs(ler("bcb_ptax_trecho.json"))
    assert len(pontos) == 51
    assert pontos[0][0] == date(2026, 7, 1)
    assert all(4 < v < 7 for _, v in pontos)


def test_sgs_que_responde_objeto_no_lugar_da_serie_e_recusa():
    with pytest.raises(fontes.FormatoInesperado):
        fontes.ler_sgs(b'{"error": "indisponivel"}')


# ------------------------------------------------------------------ RSS

def test_o_veiculo_sai_do_titulo_e_vai_para_a_fonte():
    itens = fontes.ler_rss(ler("gnews_diesel_trecho.xml"))
    assert len(itens) == 3
    for i in itens:
        assert i["fonte"] and not i["titulo"].endswith(" - " + i["fonte"])
        assert i["link"].startswith("https://news.google.com/")
        assert i["publicada_em"].tzinfo is not None


def test_link_que_nao_e_https_NAO_ENTRA():
    """O link vira `href` na tela. `javascript:` ou `http:` não passam."""
    texto = ler("gnews_diesel_trecho.xml").decode("utf-8")
    # o <channel> tem o PRÓPRIO <link> antes do primeiro item — a troca é
    # dentro do item, que é o que vira href
    i = texto.index("<item>")
    texto = texto[:i] + texto[i:].replace("<link>https://", "<link>javascript://", 1)
    assert len(fontes.ler_rss(texto.encode("utf-8"))) == 2


def test_rss_que_nao_e_xml_e_recusa():
    with pytest.raises(fontes.FormatoInesperado):
        fontes.ler_rss(b"<html><body>captcha</body>")


# ------------------------------------------------------------- a falha dita

def test_a_falha_e_dita_pelo_TIPO_e_nao_despeja_a_url():
    exc = urllib.error.HTTPError(fontes.url_noticias("diesel"), 503,
                                 "Service Unavailable", {}, None)
    assert fontes.descrever_falha(exc) == "HTTPError 503"


def test_todo_tema_corta_o_velho_na_origem():
    for tema, cfg in fontes.TEMAS.items():
        assert "when:" in cfg["busca"], tema
        assert cfg["rotulo"], tema
