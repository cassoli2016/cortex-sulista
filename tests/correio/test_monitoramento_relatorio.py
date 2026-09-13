# -*- coding: utf-8 -*-
"""O e-mail de monitoramento de cliente: o recorte do dia e as contas da
planilha da torre, contra linhas no formato REAL do `AGORA_SQL`.

As linhas abaixo são literais, com a forma do que o ERP devolveu em
12/09/2026 (datas em "YYYY-MM-DD HH:MM", CVA grudado na ordem de frete do
cliente em `ref_cliente`) — não derivadas do código que vai lê-las. Nenhum
teste vai ao ERP, ao SMTP ou à rede.
"""
from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest

from api.correio import monitoramento as mo
from api.correio import planilha_monitoramento as pm

AGORA = datetime(2026, 9, 11, 14, 0)
CPF_MOTORISTA = "31072584824"

VAZIO = {k: None for k in (
    "t_cheg_carga", "t_aguard_carga", "t_saiu_carga", "t_viagem", "t_cheg_desc",
    "t_aguard_desc", "t_fim_desc", "t_finalizada", "mdfe_em",
    "lat_origem", "lon_origem", "lat_destino", "lon_destino")}


def _r(**kw):
    kw.setdefault("coleta_chave", "1|1|2|1|0|1|%s" % kw.get("coleta"))
    base = {**VAZIO, "emissao": "2026-09-10", "origem": "CRUZEIRO", "uf_origem": "SP",
            "destino": "RESENDE", "uf_destino": "RJ", "mdfe_encerrado": 0,
            "mdfe_autorizado": 0, "carreta": "", "ref_cliente": "6100416533"}
    return {**base, **kw}


LINHAS = [
    # A — escadas, descarregou hoje DENTRO da carência de 6h30
    _r(coleta=20271, placa="BCW8A71", carreta="AQJ9G39", mercadoria="ESCADAS",
       destinatario_nome="VOLKSWAGEN - RESENDE/RJ",
       ref_cliente="6100416533CIF0103984", motorista=CPF_MOTORISTA,
       motorista_nome="JOEL DA SILVA SANTOS",
       janela_carga="2026-09-10 15:00", janela_entrega="2026-09-11 08:00",
       t_cheg_carga="2026-09-10 14:02", t_saiu_carga="2026-09-10 16:40",
       t_viagem="2026-09-10 16:45", t_cheg_desc="2026-09-11 07:45",
       t_fim_desc="2026-09-11 11:32", mdfe_encerrado=1, mdfe_em="2026-09-11 07:47"),
    # B — conjunto phevus: genérica de 3h (a equivalência com conjuntos, de
    # 11/09/2026, foi revertida em 13/09 — o relatório SAC do ERP dá a genérica)
    _r(coleta=20250, placa="JJH4J10", carreta="AQJ9632", mercadoria="CONJUNTO PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-10 23:00", janela_entrega="2026-09-11 06:30",
       t_saiu_carga="2026-09-11 00:10", t_cheg_desc="2026-09-11 06:52",
       t_fim_desc="2026-09-11 10:23"),
    # C — longarina phevus: genérica de 3h, ficou 6h12 → 3h12 parado
    _r(coleta=20274, placa="FWT9B55", carreta="AMN5F90", mercadoria="LONGARINA PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-10 21:00", janela_entrega="2026-09-11 06:00",
       t_saiu_carga="2026-09-10 22:00", t_cheg_desc="2026-09-11 07:28",
       t_fim_desc="2026-09-11 13:40"),
    # D — chegou pelo MANIFESTO e o fim da descarga não foi registrado
    _r(coleta=11930, placa="NYP3J22", carreta="AQJ9G34", mercadoria="EMBALAGENS",
       destinatario_nome="IOCHPE MAXION - CRUZEIRO/SP", destino="CRUZEIRO",
       uf_destino="SP", janela_carga="2026-09-10 15:00",
       janela_entrega="2026-09-11 10:00", mdfe_encerrado=1, mdfe_em="2026-09-11 08:13"),
    # E — em viagem, previsão pelo histórico JÁ PASSOU (não pode sair)
    _r(coleta=20282, placa="BCZ4A85", mercadoria="LONGARINA PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-11 08:00", janela_entrega="2026-09-11 15:30",
       t_saiu_carga="2026-09-11 10:00", t_viagem="2026-09-11 10:05"),
    # E2 — em viagem, previsão no FUTURO (sai)
    _r(coleta=20283, placa="FYW2E04", mercadoria="CONJUNTO PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-11 10:00", janela_entrega="2026-09-11 18:00",
       t_saiu_carga="2026-09-11 13:30"),
    # F — programada para daqui a dois dias: fora do horizonte de 24 h
    _r(coleta=20290, placa="GIY8H66", mercadoria="LONGARINA PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-13 10:00", janela_entrega="2026-09-13 14:00"),
    # G — descarregou ONTEM: é da aba de ontem
    _r(coleta=20200, placa="FOS2D32", mercadoria="LONGARINA PHEVUS",
       destinatario_nome="IOCHPE MAXION - RESENDE/RJ",
       janela_carga="2026-09-09 21:00", janela_entrega="2026-09-10 06:00",
       t_cheg_desc="2026-09-10 07:00", t_fim_desc="2026-09-10 20:00"),
    # H — destinatário com marcação no nome (o e-mail tem de escapar)
    _r(coleta=20299, placa="QCP8A94", mercadoria="RODAS",
       destinatario_nome="<script>alert(1)</script>",
       janela_carga="2026-09-11 16:00", janela_entrega="2026-09-11 22:00"),
]

# A ponte coleta -> CT-e: a 20271 tem DOIS CT-es (o mesmo veículo, números
# seguidos — vale o mais novo), a 11930 tem um, e as outras ainda nenhum.
PONTE = [
    {"coleta_chave": "1|1|2|1|0|1|20271", "grupo": 1, "empresa": 1, "filial": 2,
     "numero": 102629, "serie": 1},
    {"coleta_chave": "1|1|2|1|0|1|20271", "grupo": 1, "empresa": 1, "filial": 2,
     "numero": 102630, "serie": 1},
    {"coleta_chave": "1|1|2|1|0|1|11930", "grupo": 1, "empresa": 1, "filial": 3,
     "numero": 55120, "serie": 1},
    # coleta de OUTRO dia, fora do e-mail: não pode virar link de ninguém
    {"coleta_chave": "1|1|2|1|0|1|20200", "grupo": 1, "empresa": 1, "filial": 2,
     "numero": 102001, "serie": 1},
]

CONTRATO = [
    {"mercadoria": "", "ft_carga_h": 3.0, "ft_descarga_h": 3.0},
    {"mercadoria": "CONJUNTOS", "ft_carga_h": 3.0, "ft_descarga_h": 6.5},
    {"mercadoria": "ESCADAS", "ft_carga_h": 3.0, "ft_descarga_h": 6.5},
    {"mercadoria": "RODAS", "ft_carga_h": 3.0, "ft_descarga_h": 6.5},
]


@pytest.fixture
def erp(monkeypatch):
    """Dublê do ERP inteiro que `dados()` toca. Devolve a lista de linhas,
    que o teste pode trocar."""
    from api import db, frota_identidade, portal_cliente as pc, raster_eventos
    estado = {"linhas": list(LINHAS), "falha": None, "ponte": list(PONTE),
              "ponte_falha": None}

    def query(sql, params=None):
        assert params and params["raiz"] == "12345678"
        if "ponte coleta -> CT-e" in sql:
            if estado["ponte_falha"]:
                raise estado["ponte_falha"]
            return [dict(r) for r in estado["ponte"]]
        if estado["falha"]:
            raise estado["falha"]
        return [dict(r) for r in estado["linhas"]]

    monkeypatch.setattr(db, "query", query)
    monkeypatch.setattr(pc, "_freetime", lambda raiz: {"linhas": CONTRATO})
    monkeypatch.setattr(pc, "nome_do_cliente", lambda raiz: "CLIENTE TESTE S.A.")
    monkeypatch.setattr(pc, "_eta_por_rota",
                        lambda: {"CRUZEIRO/SP|RESENDE/RJ": (2.0, 40)})
    monkeypatch.setattr(raster_eventos, "por_placa",
                        lambda *a, **k: {"macros": {}, "trilhas": {}})
    monkeypatch.setattr(frota_identidade, "mapa", lambda *a: {
        "BCW8A71": {"frota": "T3008"}, "AQJ9G39": {"frota": "G3006"},
        # placa copiada no campo frota NÃO é número de frota
        "JJH4J10": {"frota": "JJH4J10"}})
    return estado


def _d():
    return mo.dados("12345678", agora=AGORA)


def _por(d):
    return {c["coleta"]: c for c in d["cargas"]}


def test_o_dia_e_o_da_aba_da_torre(erp):
    """No ar, programada para as próximas 24 h, e descarregada HOJE."""
    assert set(_por(_d())) == {20271, 20250, 20274, 11930, 20282, 20283, 20299}


def test_carencia_e_a_do_contrato_pela_mercadoria(erp):
    c = _por(_d())
    assert c[20271]["carencia_h"] == 6.5                 # cláusula de escadas
    assert c[20250]["carencia_h"] == 3.0                 # genérica (equivalência revertida)
    assert c[20274]["carencia_h"] == 3.0                 # genérica


def test_horas_paradas_so_com_as_DUAS_pontas_registradas(erp):
    """Relógio correndo sobre um fim de descarga que atrasa mede o atraso do
    registro — e na caixa do cliente vira cobrança de espera de um veículo
    que já foi embora (visto no dado real de 12/09/2026)."""
    c = _por(_d())
    assert c[20271]["paradas_h"] == 0
    assert c[20274]["paradas_h"] == pytest.approx(3.2, abs=0.01)
    assert c[11930]["situacao"] == "no_cliente"
    assert c[11930]["chegada"] == "2026-09-11 08:13"
    assert c[11930]["permanencia_h"] is None and c[11930]["paradas_h"] is None


def test_previsao_vencida_nao_sai(erp):
    c = _por(_d())
    assert c[20282]["eta"] is None
    assert c[20283]["eta"] == "2026-09-11 15:30"


def test_cva_so_quando_o_erp_tem(erp):
    """Onde o ERP não tem o CVA a coluna fica em branco: fabricar "CIF00" +
    pedido, como a torre faz à mão, seria inventar um documento."""
    c = _por(_d())
    assert c[20271]["cva"] == "CIF0103984"
    assert c[20250]["cva"] == ""


def test_frota_e_placas_como_na_planilha(erp):
    c = _por(_d())
    assert c[20271]["frota"] == "T3008 / G3006"
    assert c[20271]["placas"] == "BCW8A71 / AQJ9G39"
    assert c[20250]["frota"] == ""                       # placa no campo frota


def test_grupos_em_ordem_de_dicionario(erp):
    """A mensagem chega de duas em duas horas: quem lê aprende onde fica o
    seu destino, e ordem que muda a cada envio obriga a procurar toda vez."""
    g = [(x["mercadoria"], x["destinatario"]) for x in _d()["grupos"]]
    assert g == sorted(g)


def test_motorista_so_pelo_PRIMEIRO_nome_e_o_CPF_nunca(erp):
    """Decisão de quem opera (12/09/2026): o primeiro nome vai, como sempre
    foi na planilha da torre. Sobrenome não, e o CPF — que é o próprio código
    do motorista na coleta — em lugar nenhum."""
    r = mo.montar("12345678", agora=AGORA)
    wb = openpyxl.load_workbook(io.BytesIO(r["anexos"][0]["conteudo"]))
    valores = [str(v) for row in wb.active.iter_rows(values_only=True) for v in row if v]
    assert "Joel" in r["html"] and "Joel" in r["texto"] and "Joel" in valores
    for proibido in (CPF_MOTORISTA, "SILVA SANTOS", "Silva"):
        assert proibido not in r["html"] and proibido not in r["texto"]
        assert not any(proibido in v for v in valores)


def test_html_escapa_o_que_vem_do_erp(erp):
    r = mo.montar("12345678", agora=AGORA)
    assert "<script>alert(1)</script>" not in r["html"]
    assert "&lt;script&gt;" in r["html"]


def test_e_mail_diz_o_que_se_sabe_e_o_que_falta(erp):
    r = mo.montar("12345678", agora=AGORA)
    assert "fim da descarga não registrado" in r["html"]
    assert "+3h12 parado" in r["html"]
    assert "dentro da carência" in r["html"]
    assert r["assunto"].startswith("Monitoramento de cargas · CLIENTE TESTE")
    assert "1 no cliente" in r["assunto"]


def test_sem_carga_no_dia_e_VAZIO_e_nao_anexa(erp):
    erp["linhas"] = [LINHAS[6], LINHAS[7]]               # F e G: nada do dia
    r = mo.montar("12345678", agora=AGORA)
    assert r["vazio"] is True and r["anexos"] == []


def test_raiz_invalida_e_recusada_antes_do_erp(erp):
    with pytest.raises(ValueError):
        mo.dados("", agora=AGORA)


# ─────────────────────────────────────────────────────────── a planilha ────

def test_planilha_no_formato_da_torre(erp):
    wb = openpyxl.load_workbook(io.BytesIO(pm.gerar(_d())))
    ws = wb.active
    assert ws.title == "11.09"
    assert [c.value for c in ws[1]] == [n for n, _ in pm.COLUNAS]
    assert ws.freeze_panes == "A2"
    assert ws["A2"].value and "→" in ws["A2"].value     # faixa do fluxo
    assert any(str(m) == "A2:N2" for m in ws.merged_cells.ranges)
    linha = next(r for r in ws.iter_rows(min_row=2) if r[10].value == 20271)
    assert linha[2].value == "Joel"
    total, carencia, paradas = linha[7], linha[8], linha[9]
    assert total.number_format == "[h]:mm" and total.value.total_seconds() == 3 * 3600 + 47 * 60
    assert carencia.value.total_seconds() == 6.5 * 3600
    assert paradas.value.total_seconds() == 0
    assert linha[11].value == "CIF0103984"


def test_planilha_deixa_em_branco_o_que_nao_foi_registrado(erp):
    """Como a da torre: sem a saída, total e horas paradas ficam vazios."""
    ws = openpyxl.load_workbook(io.BytesIO(pm.gerar(_d()))).active
    linha = next(r for r in ws.iter_rows(min_row=2) if r[10].value == 11930)
    assert linha[5].value is not None                   # chegada
    assert linha[6].value is None and linha[7].value is None and linha[9].value is None
    assert "não registrado" in linha[13].value


# ─────────────────────────────────────────────────────────── a rodada ──────

def _mon(**kw):
    return {"id": 9, "cliente_raiz": "12345678", "cliente_nome": "CLIENTE TESTE",
            "mercadorias": [], "destinatarios": "logistica@cliente.test",
            "responder_para": "torre@sulista.test", "anexar_planilha": True, **kw}


def test_rodada_manda_com_planilha_e_resposta_para_a_torre(erp):
    enviados, passagens = [], []

    def enviar(dest, assunto, texto, **kw):
        enviados.append({"dest": dest, "assunto": assunto, **kw})
        return {"ok": True, "erro": ""}

    linha = mo.rodar(_mon(), enviar=enviar,
                     montar_=lambda raiz, mercs, anexar: mo.montar(
                         raiz, mercs, anexar=anexar, agora=AGORA),
                     registrar=lambda i, r: passagens.append((i, r)))
    assert linha.startswith("OK")
    assert enviados[0]["dest"] == "logistica@cliente.test"
    assert enviados[0]["responder_para"] == "torre@sulista.test"
    assert enviados[0]["anexos"][0]["nome"].endswith(".xlsx")
    assert passagens == [(9, "enviado")]


def test_ERP_fora_NADA_vai_ao_cliente_e_a_passagem_e_falha(erp):
    """Relatório interno manda "não consegui ler"; este vai para o CLIENTE,
    e mensagem de falha do nosso sistema na caixa dele não serve a ninguém."""
    erp["falha"] = RuntimeError("timeout do ERP")
    enviados, passagens = [], []
    linha = mo.rodar(_mon(), enviar=lambda *a, **k: enviados.append(a),
                     registrar=lambda i, r: passagens.append(r))
    assert enviados == []
    assert linha.startswith("FALHA") and "nada enviado" in linha
    assert passagens and passagens[0].startswith("falhou")
    # o texto da exceção (que carrega trecho de consulta) não vai para a trilha
    assert "timeout do ERP" not in passagens[0]


def test_dia_sem_carga_CALA_e_marca_a_passagem(erp):
    erp["linhas"] = [LINHAS[7]]
    enviados, passagens = [], []
    linha = mo.rodar(_mon(), enviar=lambda *a, **k: enviados.append(a),
                     montar_=lambda raiz, mercs, anexar: mo.montar(
                         raiz, mercs, anexar=anexar, agora=AGORA),
                     registrar=lambda i, r: passagens.append(r))
    assert enviados == [] and "sem carga" in linha
    assert passagens == ["sem carga no dia — não enviado"]


def test_ensaio_nao_envia_nem_marca(erp):
    enviados, passagens = [], []
    linha = mo.rodar(_mon(), ensaio=True, enviar=lambda *a, **k: enviados.append(a),
                     montar_=lambda raiz, mercs, anexar: mo.montar(
                         raiz, mercs, anexar=anexar, agora=AGORA),
                     registrar=lambda i, r: passagens.append(r))
    assert enviados == [] and passagens == [] and "enviaria" in linha


# ─────────────────────────────────────────────── o envelope do e-mail ──────

def test_xlsx_vai_com_o_tipo_certo_e_a_resposta_vai_para_a_torre():
    from api.correio import envio
    msg = envio._mensagem(["a@cliente.test"], "assunto", "texto", None,
                          {"remetente": "cortex@sulista.test"},
                          anexos=[{"nome": "m.xlsx", "conteudo": b"PK\x03\x04"}],
                          responder_para=["torre@sulista.test"])
    assert msg["Reply-To"] == "torre@sulista.test"
    anexo = [p for p in msg.iter_attachments()][0]
    assert anexo.get_content_type() == \
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_resposta_invalida_e_recusada_no_envio():
    from api.correio import envio
    r = envio.enviar("a@cliente.test", "x", "y", registrar=False,
                     responder_para="torre@")
    assert not r["ok"] and "resposta" in r["erro"]


# ─────────────────────────────────────────────── ida e volta no banco ──────

def test_gravar_listar_e_remover(esquema_pg):
    v = mo.gravar({"cliente_raiz": "12345678", "cliente_nome": "CLIENTE TESTE",
                   "destinatarios": "a@cliente.test", "mercadorias": ["ESCADAS"],
                   "dias_semana": [1, 2, 3, 4, 5, 6], "ativo": True},
                  "admin@sulista.test", esquema=esquema_pg)
    itens = mo.listar(esquema=esquema_pg)
    assert [x["id"] for x in itens] == [v["id"]]
    assert itens[0]["mercadorias"] == v["mercadorias"] and itens[0]["ativo"] is True
    mo.gravar({**v, "ativo": False}, "admin@sulista.test", esquema=esquema_pg)
    assert mo.listar(esquema=esquema_pg)[0]["ativo"] is False
    mo.registrar_execucao(v["id"], "enviado", esquema=esquema_pg)
    assert mo.listar(esquema=esquema_pg)[0]["ultimo_resultado"] == "enviado"
    mo.remover(v["id"], esquema=esquema_pg)
    assert mo.listar(esquema=esquema_pg) == []


# ───────────────────────────────────────── o link "Ver onde está a carga" ──

def _abre(link):
    from api import url_publica
    from api.rastreio import consulta
    assert link.startswith(url_publica.base() + "/r#c="), link
    # a marca de origem vai no fim do fragmento — sem ela a página diria
    # "você já recebe por WhatsApp" a quem veio do e-mail
    assert link.endswith("&o=email"), link
    return consulta.link_abrir(link.split("#c=", 1)[1].split("&", 1)[0])


def test_o_link_abre_o_CT_e_MAIS_NOVO_da_coleta(erp):
    """Pedido de quem opera (13/09/2026). É o MESMO link assinado do aviso de
    WhatsApp, e ele tem de abrir a carga certa: dois CT-es da mesma coleta são
    o mesmo veículo, e vale o mais recente."""
    c = _por(_d())
    assert _abre(c[20271]["link"]) == {"g": 1, "e": 1, "f": 2, "n": 102630, "s": 1}
    assert _abre(c[11930]["link"]) == {"g": 1, "e": 1, "f": 3, "n": 55120, "s": 1}


def test_carga_sem_CT_e_fica_sem_link(erp):
    """Carga programada ainda não tem CT-e, e o rastreio só conhece CT-e: um
    link para a página vazia seria pior que nenhum."""
    c = _por(_d())
    assert c[20250]["link"] is None and c[20299]["link"] is None


def test_ponte_fora_do_ar_NAO_derruba_o_email(erp):
    """O link é acréscimo: sem ele o e-mail sai igual, só sem o botão."""
    erp["ponte_falha"] = RuntimeError("timeout")
    r = mo.montar("12345678", agora=AGORA)
    # o ENDEREÇO, e não a frase: o "Como ler" explica o link com as mesmas
    # palavras e continua lá mesmo quando nenhum link sai
    assert r["vazio"] is False and "/r#c=" not in r["html"]
    assert "Onde está: " not in r["texto"]
    assert all(c["link"] is None for c in _d()["cargas"])


def test_o_link_vai_no_email_no_texto_e_na_planilha(erp):
    d = _d()
    link = _por(d)[20271]["link"]
    r = mo.montar("12345678", dados_=d)
    # no atributo HTML o `&` sai escapado, que é o certo — o navegador desfaz
    assert 'href="%s"' % link.replace("&", "&amp;") in r["html"]
    assert "Ver onde está a carga" in r["html"]
    assert "Onde está: " + link in r["texto"]
    ws = openpyxl.load_workbook(io.BytesIO(r["anexos"][0]["conteudo"])).active
    linha = next(x for x in ws.iter_rows(min_row=2) if x[10].value == 20271)
    assert linha[12].value == "abrir" and linha[12].hyperlink.target == link
    sem = next(x for x in ws.iter_rows(min_row=2) if x[10].value == 20250)
    assert sem[12].value is None and sem[12].hyperlink is None
